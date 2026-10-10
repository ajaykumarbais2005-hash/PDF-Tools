import io
import tempfile
import unittest
import zipfile
from pathlib import Path

import docx
import pptx
from PIL import Image, ImageDraw, ImageFont
from reportlab.pdfgen import canvas

import backend
import auth_store
from web_app import create_app


class WebsiteFlowTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        root = Path(self.temp_dir.name)
        self.store = auth_store.AuthStore(str(root / "accounts.sqlite3"))
        self.app = create_app(self.store, root / "web_jobs")
        self.app.config.update(TESTING=True, SECRET_KEY="test-secret")
        self.client = self.app.test_client()

    def tearDown(self):
        self.temp_dir.cleanup()

    def csrf(self):
        with self.client.session_transaction() as session:
            return session["_csrf_token"]

    def create_admin(self):
        self.client.get("/setup")
        response = self.client.post(
            "/setup",
            data={
                "_csrf_token": self.csrf(),
                "username": "admin",
                "password": "long-enough-test-password",
            },
        )
        self.assertEqual(response.status_code, 302)

    @staticmethod
    def sample_pdf():
        stream = io.BytesIO()
        document = canvas.Canvas(stream)
        document.drawString(72, 720, "Website test document")
        document.save()
        stream.seek(0)
        return stream

    def test_setup_login_and_tool_download(self):
        response = self.client.get("/")
        self.assertEqual(response.status_code, 302)
        self.assertIn("/setup", response.location)
        self.create_admin()

        response = self.client.get("/")
        self.assertEqual(response.status_code, 200)
        self.assertIn(b"Find the right tool", response.data)
        icon_response = self.client.get("/brand-icon")
        self.assertEqual(icon_response.status_code, 200)
        icon_response.close()

        response = self.client.post(
            "/api/tools/pdf-information",
            headers={"X-CSRF-Token": self.csrf()},
            data={"file": (self.sample_pdf(), "source.pdf")},
            content_type="multipart/form-data",
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json["result"]["page_count"], 1)
        self.assertNotIn("file_path", response.json["result"])

        response = self.client.post(
            "/api/tools/compress-pdf",
            headers={"X-CSRF-Token": self.csrf()},
            data={
                "file": (self.sample_pdf(), "source.pdf"),
                "level": "low",
            },
            content_type="multipart/form-data",
        )
        self.assertEqual(response.status_code, 200, response.json)
        download = self.client.get(response.json["download_url"])
        self.assertEqual(download.status_code, 200)
        self.assertEqual(download.mimetype, "application/pdf")
        download.close()

    def test_admin_invitation_registration_and_account_controls(self):
        self.create_admin()
        csrf = self.csrf()
        invitation_response = self.client.post(
            "/api/admin/invitations",
            headers={"X-CSRF-Token": csrf},
            json={},
        )
        self.assertEqual(invitation_response.status_code, 200)
        invitation = invitation_response.json["invitation"]

        self.client.post("/logout", data={"_csrf_token": csrf})
        self.client.get("/register")
        response = self.client.post(
            "/register",
            data={
                "_csrf_token": self.csrf(),
                "invitation": invitation,
                "username": "member",
                "password": "another-long-test-password",
            },
        )
        self.assertEqual(response.status_code, 302)

        self.assertEqual(
            self.client.get("/api/admin/users").status_code,
            403,
        )

        self.client.post("/logout", data={"_csrf_token": self.csrf()})
        self.client.get("/login")
        admin_csrf = self.csrf()
        response = self.client.post(
            "/login",
            data={
                "_csrf_token": admin_csrf,
                "username": "admin",
                "password": "long-enough-test-password",
            },
        )
        self.assertEqual(response.status_code, 302)
        admin_csrf = self.csrf()
        response = self.client.post(
            "/api/admin/users/member/password",
            headers={"X-CSRF-Token": admin_csrf},
            json={"password": "reset-long-test-password"},
        )
        self.assertEqual(response.status_code, 200, response.json)
        response = self.client.post(
            "/api/admin/users/member/enabled",
            headers={"X-CSRF-Token": admin_csrf},
            json={"enabled": False},
        )
        self.assertEqual(response.status_code, 200, response.json)
        with self.assertRaises(auth_store.AuthenticationError):
            self.store.authenticate("member", "reset-long-test-password")

    def test_api_rejects_missing_csrf_token(self):
        self.create_admin()
        response = self.client.post("/api/tools/pdf-information")
        self.assertEqual(response.status_code, 400)
        self.assertIn("session expired", response.json["error"].lower())

    def test_all_twenty_tools_run_end_to_end(self):
        self.create_admin()
        csrf = self.csrf()

        pdf_stream = io.BytesIO()
        pdf_canvas = canvas.Canvas(pdf_stream)
        for page_number in range(1, 4):
            pdf_canvas.drawString(72, 720, f"PDF Tools test page {page_number}")
            pdf_canvas.showPage()
        pdf_canvas.save()
        pdf_bytes = pdf_stream.getvalue()

        word_stream = io.BytesIO()
        word_document = docx.Document()
        word_document.add_heading("Test document", level=1)
        word_document.add_paragraph("PDF Tools conversion test.")
        word_document.save(word_stream)
        word_bytes = word_stream.getvalue()

        presentation_stream = io.BytesIO()
        presentation = pptx.Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = "PDF Tools slide"
        presentation.save(presentation_stream)
        presentation_bytes = presentation_stream.getvalue()

        def run_tool(name, fields=None, uploads=None):
            data = dict(fields or {})
            if uploads:
                data.update(uploads)
            response = self.client.post(
                f"/api/tools/{name}",
                headers={"X-CSRF-Token": csrf},
                data=data,
                content_type="multipart/form-data",
            )
            self.assertEqual(response.status_code, 200, f"{name}: {response.get_data(as_text=True)}")
            return response.json

        def download_result(payload, extension=None):
            self.assertIn("download_url", payload)
            response = self.client.get(payload["download_url"])
            self.assertEqual(response.status_code, 200)
            content = response.get_data()
            disposition = response.headers["Content-Disposition"].lower()
            response.close()
            self.assertGreater(len(content), 0)
            if extension:
                self.assertTrue(
                    disposition.endswith(extension.lower()),
                    disposition,
                )
            return content

        payload = run_tool("create-pdf", {"title": "Created PDF", "text_content": "Created from the web app."})
        download_result(payload, ".pdf")

        payload = run_tool("merge-pdfs", uploads={
            "files": [(io.BytesIO(pdf_bytes), "first.pdf"), (io.BytesIO(pdf_bytes), "second.pdf")]
        })
        self.assertEqual(payload["result"]["total_pages"], 6)
        download_result(payload, ".pdf")

        payload = run_tool("split-pdf", uploads={"file": (io.BytesIO(pdf_bytes), "source.pdf")})
        split_bytes = download_result(payload, ".zip")
        with zipfile.ZipFile(io.BytesIO(split_bytes)) as archive:
            self.assertEqual(len(archive.namelist()), 3)

        payload = run_tool("lock-pdf", {"user_password": "test-password"},
                           {"file": (io.BytesIO(pdf_bytes), "source.pdf")})
        locked_pdf = download_result(payload, ".pdf")
        payload = run_tool("unlock-pdf", {"password": "test-password"},
                           {"file": (io.BytesIO(locked_pdf), "locked.pdf")})
        download_result(payload, ".pdf")

        payload = run_tool("compress-pdf", uploads={"file": (io.BytesIO(pdf_bytes), "source.pdf")})
        download_result(payload, ".pdf")

        payload = run_tool("watermark-pdf", uploads={"file": (io.BytesIO(pdf_bytes), "source.pdf")})
        self.assertEqual(payload["result"]["pages_watermarked"], 3)
        download_result(payload, ".pdf")

        payload = run_tool("pdf-to-ppt", uploads={"file": (io.BytesIO(pdf_bytes), "source.pdf")})
        ppt_bytes = download_result(payload, ".pptx")
        self.assertEqual(payload["result"]["ocr_pages"], 0)
        converted_ppt = pptx.Presentation(io.BytesIO(ppt_bytes))
        converted_text = "\n".join(
            shape.text
            for slide in converted_ppt.slides
            for shape in slide.shapes
            if shape.has_text_frame
        )
        self.assertIn("PDF Tools test page 1", converted_text)
        invalid_dpi = self.client.post(
            "/api/tools/pdf-to-ppt",
            headers={"X-CSRF-Token": csrf},
            data={
                "dpi": "601",
                "file": (io.BytesIO(pdf_bytes), "source.pdf"),
            },
            content_type="multipart/form-data",
        )
        self.assertEqual(invalid_dpi.status_code, 400)
        self.assertIn("36 and 600 DPI", invalid_dpi.json["error"])

        payload = run_tool("pdf-to-word", uploads={"file": (io.BytesIO(pdf_bytes), "source.pdf")})
        download_result(payload, ".docx")

        payload = run_tool("pdf-to-image", uploads={"file": (io.BytesIO(pdf_bytes), "source.pdf")})
        images_zip = download_result(payload, ".zip")
        with zipfile.ZipFile(io.BytesIO(images_zip)) as archive:
            self.assertEqual(len(archive.namelist()), 3)

        payload = run_tool("ocr", uploads={"file": (io.BytesIO(pdf_bytes), "source.pdf")})
        extracted_ocr = download_result(payload, ".txt").decode("utf-8")
        self.assertIn("test page 1", extracted_ocr)

        payload = run_tool("extract-text", uploads={"file": (io.BytesIO(pdf_bytes), "source.pdf")})
        extracted_text = download_result(payload, ".txt").decode("utf-8")
        self.assertIn("test page 3", extracted_text)

        payload = run_tool("ppt-to-pdf", uploads={"file": (io.BytesIO(presentation_bytes), "source.pptx")})
        download_result(payload, ".pdf")

        payload = run_tool("word-to-pdf", uploads={"file": (io.BytesIO(word_bytes), "source.docx")})
        download_result(payload, ".pdf")

        payload = run_tool("rotate-pdf", uploads={"file": (io.BytesIO(pdf_bytes), "source.pdf")})
        self.assertEqual(payload["result"]["pages_rotated"], 3)
        download_result(payload, ".pdf")

        payload = run_tool("delete-pages", {"pages_to_delete": "2"},
                           {"file": (io.BytesIO(pdf_bytes), "source.pdf")})
        self.assertEqual(payload["result"]["pages_remaining"], 2)
        download_result(payload, ".pdf")

        payload = run_tool("reorder-pages", {"page_order": "3, 2, 1"},
                           {"file": (io.BytesIO(pdf_bytes), "source.pdf")})
        download_result(payload, ".pdf")

        payload = run_tool("pdf-information", uploads={"file": (io.BytesIO(pdf_bytes), "source.pdf")})
        self.assertEqual(payload["result"]["page_count"], 3)

        payload = run_tool("update-metadata", {"title": "Updated title", "author": "Test author"},
                           {"file": (io.BytesIO(pdf_bytes), "source.pdf")})
        download_result(payload, ".pdf")

        payload = run_tool("pdf-preview", {"page_number": "2", "zoom": "1"},
                           {"file": (io.BytesIO(pdf_bytes), "source.pdf")})
        self.assertEqual(payload["result"]["page_number"], 2)
        self.assertTrue(payload["result"]["image"].startswith("data:image/png;base64,"))

    def test_pdf_to_ppt_uses_ocr_for_scanned_text(self):
        if not backend._ensure_tesseract_on_path():
            self.skipTest("Tesseract OCR is not installed.")
        if not backend._find_tessdata_dir("eng"):
            self.skipTest("Tesseract English language data is not installed.")
        self.create_admin()

        image = Image.new("RGB", (1600, 400), "white")
        draw = ImageDraw.Draw(image)
        draw.text((80, 120), "Editable scanned text", fill="black",
                  font=ImageFont.truetype(r"C:\Windows\Fonts\arial.ttf", 100))
        scanned_pdf = io.BytesIO()
        pdf_canvas = canvas.Canvas(scanned_pdf, pagesize=(612, 792))
        pdf_canvas.drawInlineImage(image, 36, 320, width=540, height=135)
        pdf_canvas.save()

        payload = self.client.post(
            "/api/tools/pdf-to-ppt",
            headers={"X-CSRF-Token": self.csrf()},
            data={"file": (io.BytesIO(scanned_pdf.getvalue()), "scanned.pdf")},
            content_type="multipart/form-data",
        )
        self.assertEqual(payload.status_code, 200, payload.get_data(as_text=True))
        self.assertEqual(payload.json["result"]["ocr_pages"], 1)
        download = self.client.get(payload.json["download_url"])
        converted_bytes = download.get_data()
        download.close()
        converted_ppt = pptx.Presentation(io.BytesIO(converted_bytes))
        converted_text = "\n".join(
            shape.text
            for slide in converted_ppt.slides
            for shape in slide.shapes
            if shape.has_text_frame
        )
        self.assertIn("editable scanned text", " ".join(converted_text.casefold().split()))


if __name__ == "__main__":
    unittest.main()
